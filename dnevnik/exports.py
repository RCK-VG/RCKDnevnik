"""Builds the Excel/CSV export files: attendance matrix, absence summary,
notes, and (admin-only) a full data dump. Kept separate from views_export.py
so the file-building logic can be unit tested without going through HTTP.
"""

import csv
import io
from urllib.parse import quote

import openpyxl
from django.http import HttpResponse
from django.utils import timezone
from django.utils.timezone import localtime

from . import constants
from .models import Attendance, ExportLog, Lesson, Note, SchoolClass, SchoolYear, Student, Subject
from .stats import stats_aggregate_kwargs

CSV_DELIMITER = ";"

ATTENDANCE_LABELS = {
    constants.ATTENDANCE_PRESENT: "Prisutan",
    constants.ATTENDANCE_LATE: "Kasni",
}

# Standard Excel "Good/Bad/Neutral" cell-style colors, reused here so the
# exported matrix looks native to anyone who has used Excel's own
# conditional formatting.
FILL_PRESENT, FONT_PRESENT = "C6EFCE", "006100"
FILL_LATE, FONT_LATE = "FFEB9C", "9C6500"
FILL_ABSENT, FONT_ABSENT = "FFC7CE", "9C0006"


def _attendance_cell(attendance):
    if attendance is None:
        return ""
    if attendance.status == constants.ATTENDANCE_ABSENT:
        if attendance.justified is True:
            return "Odsutan (opravdano)"
        if attendance.justified is False:
            return "Odsutan (neopravdano)"
        return "Odsutan"
    return ATTENDANCE_LABELS.get(attendance.status, attendance.status)


def _status_fill(cell_text):
    if cell_text.startswith("Prisutan"):
        return FILL_PRESENT, FONT_PRESENT
    if cell_text.startswith("Kasni"):
        return FILL_LATE, FONT_LATE
    if cell_text.startswith("Odsutan"):
        return FILL_ABSENT, FONT_ABSENT
    return None, None


def _safe_filename_part(text):
    return str(text).replace(" ", "_")


def _display_name(user):
    if user is None:
        return ""
    return user.get_full_name() or user.username


def _meta_lines(user, title):
    now = localtime(timezone.now())
    return [
        title,
        f"Izvoz izradio/la: {_display_name(user)}",
        f"Datum i vrijeme izvoza: {now.strftime('%d.%m.%Y. %H:%M')}",
    ]


def _log_export(user, report_type, filters_summary):
    ExportLog.objects.create(user=user, report_type=report_type, filters_summary=filters_summary)


# ---------------------------------------------------------------------------
# a) Attendance matrix (students x lessons) for one class, one subject or ALL
#    subjects, over a period - plus (in Excel) a Satovi and a Bilješke sheet,
#    so a razrednik gets everything for that period in one file.
# ---------------------------------------------------------------------------


def attendance_matrix_data(school_class, subject, datum_od, datum_do):
    lessons_qs = Lesson.objects.filter(school_class=school_class)
    if subject:
        lessons_qs = lessons_qs.filter(subject=subject)
    if datum_od:
        lessons_qs = lessons_qs.filter(date__gte=datum_od)
    if datum_do:
        lessons_qs = lessons_qs.filter(date__lte=datum_do)
    lessons = list(
        lessons_qs.select_related("subject").order_by("subject__name", "date", "created_at")
    )

    # Columns are (subject, date) pairs, in that order - so with "svi
    # predmeti" the table naturally groups all of one subject's dates
    # together before moving to the next subject.
    columns = []
    lessons_by_column = {}
    for lesson in lessons:
        key = (lesson.subject_id, lesson.date)
        if key not in lessons_by_column:
            columns.append((lesson.subject, lesson.date))
        lessons_by_column.setdefault(key, []).append(lesson)

    students = Student.objects.filter(school_class=school_class, is_archived=False)

    attendance_by_key = {
        (a.student_id, a.lesson_id): a
        for a in Attendance.objects.filter(lesson__in=lessons).select_related("lesson")
    }

    show_subject_in_header = subject is None
    header = ["Učenik"] + [
        (
            f"{subj.name}\n{d.strftime('%d.%m.%Y.')}"
            if show_subject_in_header
            else d.strftime("%d.%m.%Y.")
        )
        for subj, d in columns
    ]

    rows = []
    for student in students:
        row = [student.full_name]
        for subj, d in columns:
            # If more than one lesson lands on the same subject+date (e.g. a
            # covering teacher), the most recently created one wins.
            cell = ""
            for lesson in lessons_by_column[(subj.id, d)]:
                attendance = attendance_by_key.get((student.id, lesson.id))
                if attendance is not None:
                    cell = _attendance_cell(attendance)
            row.append(cell)
        rows.append(row)

    topics = [
        (subj, d, "; ".join(l.topic for l in lessons_by_column[(subj.id, d)] if l.topic))
        for subj, d in columns
    ]
    return header, rows, topics


def respond_attendance_matrix(user, school_class, subject, datum_od, datum_do, fmt):
    header, rows, topics = attendance_matrix_data(school_class, subject, datum_od, datum_do)
    subject_label = subject.name if subject else "svi predmeti"
    title = f"Prisutnost - {school_class} - {subject_label} ({datum_od or '...'} - {datum_do or '...'})"
    filters_summary = f"razred={school_class}; predmet={subject_label}; od={datum_od}; do={datum_do}"
    _log_export(user, "prisutnost", filters_summary)
    filename_base = f"prisutnost_{_safe_filename_part(school_class)}_{_safe_filename_part(subject_label)}"

    if fmt == "csv":
        content = _to_csv(_meta_lines(user, title), header, rows)
        return _attachment_response(content, f"{filename_base}.csv", "text/csv; charset=utf-8")

    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Prisutnost"
    _write_xlsx_sheet(ws, _meta_lines(user, title), header, rows, color_status_cells=True)

    ws2 = workbook.create_sheet("Satovi")
    _write_xlsx_sheet(
        ws2,
        [f"Popis satova - {school_class} - {subject_label}"],
        ["Predmet", "Datum", "Tema sata"],
        [[subj.name, d.strftime("%d.%m.%Y."), topic] for subj, d, topic in topics],
    )

    notes_header, notes_rows = notes_data(school_class, subject, None, datum_od, datum_do)
    ws3 = workbook.create_sheet("Bilješke")
    _write_xlsx_sheet(
        ws3, [f"Bilješke - {school_class} - {subject_label}"], notes_header, notes_rows
    )

    return _workbook_response(workbook, f"{filename_base}.xlsx")


# ---------------------------------------------------------------------------
# b) Absence summary per student
# ---------------------------------------------------------------------------


def absence_summary_data(school_class, subject, teacher, datum_od, datum_do):
    attendance_qs = Attendance.objects.all()
    if school_class:
        attendance_qs = attendance_qs.filter(student__school_class=school_class)
    if subject:
        attendance_qs = attendance_qs.filter(lesson__subject=subject)
    if teacher:
        attendance_qs = attendance_qs.filter(lesson__teacher=teacher)
    if datum_od:
        attendance_qs = attendance_qs.filter(lesson__date__gte=datum_od)
    if datum_do:
        attendance_qs = attendance_qs.filter(lesson__date__lte=datum_do)

    stats = attendance_qs.values(
        "student_id", "student__first_name", "student__last_name", "student__school_class__name"
    ).annotate(**stats_aggregate_kwargs())
    stats = sorted(stats, key=lambda r: (r["student__last_name"], r["student__first_name"]))

    header = ["Učenik", "Razred", "Odsutan", "Opravdano", "Neopravdano", "Kasni"]
    rows = [
        [
            f"{r['student__last_name']} {r['student__first_name']}",
            r["student__school_class__name"],
            r["odsutan"],
            r["opravdano"],
            r["neopravdano"],
            r["kasni"],
        ]
        for r in stats
    ]
    return header, rows


def respond_absence_summary(user, school_class, subject, teacher, datum_od, datum_do, fmt):
    header, rows = absence_summary_data(school_class, subject, teacher, datum_od, datum_do)
    title = "Sažetak izostanaka"
    filters_summary = (
        f"razred={school_class}; predmet={subject}; nastavnik={teacher}; "
        f"od={datum_od}; do={datum_do}"
    )
    _log_export(user, "sazetak_izostanaka", filters_summary)

    if fmt == "csv":
        content = _to_csv(_meta_lines(user, title), header, rows)
        return _attachment_response(content, "sazetak_izostanaka.csv", "text/csv; charset=utf-8")

    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Sažetak"
    _write_xlsx_sheet(ws, _meta_lines(user, title), header, rows)
    return _workbook_response(workbook, "sazetak_izostanaka.xlsx")


# ---------------------------------------------------------------------------
# c) Notes export
# ---------------------------------------------------------------------------


def notes_data(school_class, subject, teacher, datum_od, datum_do):
    notes_qs = Note.objects.filter(is_archived=False).select_related(
        "lesson__school_class", "lesson__subject", "student", "author"
    )
    if school_class:
        notes_qs = notes_qs.filter(lesson__school_class=school_class)
    if subject:
        notes_qs = notes_qs.filter(lesson__subject=subject)
    if teacher:
        notes_qs = notes_qs.filter(lesson__teacher=teacher)
    if datum_od:
        notes_qs = notes_qs.filter(lesson__date__gte=datum_od)
    if datum_do:
        notes_qs = notes_qs.filter(lesson__date__lte=datum_do)
    notes_qs = notes_qs.order_by("lesson__date", "created_at")

    header = ["Datum", "Razred", "Predmet", "Tema sata", "Učenik", "Tekst", "Autor", "Vrijeme unosa"]
    rows = [
        [
            n.lesson.date.strftime("%d.%m.%Y."),
            n.lesson.school_class.name,
            n.lesson.subject.name,
            n.lesson.topic,
            n.student.full_name if n.student else "(opća bilješka)",
            n.text,
            _display_name(n.author),
            localtime(n.created_at).strftime("%d.%m.%Y. %H:%M"),
        ]
        for n in notes_qs
    ]
    return header, rows


def respond_notes_export(user, school_class, subject, teacher, datum_od, datum_do, fmt):
    header, rows = notes_data(school_class, subject, teacher, datum_od, datum_do)
    title = "Bilješke"
    filters_summary = (
        f"razred={school_class}; predmet={subject}; nastavnik={teacher}; "
        f"od={datum_od}; do={datum_do}"
    )
    _log_export(user, "biljeske", filters_summary)

    if fmt == "csv":
        content = _to_csv(_meta_lines(user, title), header, rows)
        return _attachment_response(content, "biljeske.csv", "text/csv; charset=utf-8")

    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Bilješke"
    _write_xlsx_sheet(ws, _meta_lines(user, title), header, rows)
    return _workbook_response(workbook, "biljeske.xlsx")


# ---------------------------------------------------------------------------
# d) Full data export (admin only)
# ---------------------------------------------------------------------------


def respond_full_export(user):
    _log_export(user, "potpuni_izvoz", "")
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    ws = workbook.create_sheet("Školske godine")
    _write_xlsx_sheet(
        ws,
        _meta_lines(user, "Potpuni izvoz - Školske godine"),
        ["Naziv", "Aktivna", "Arhivirana"],
        [[y.name, y.is_active, y.is_archived] for y in SchoolYear.objects.all()],
    )

    ws = workbook.create_sheet("Razredi")
    _write_xlsx_sheet(
        ws,
        _meta_lines(user, "Potpuni izvoz - Razredi"),
        ["Naziv", "Školska godina", "Razrednik"],
        [
            [c.name, c.school_year.name, _display_name(c.homeroom_teacher)]
            for c in SchoolClass.objects.select_related("school_year", "homeroom_teacher")
        ],
    )

    ws = workbook.create_sheet("Predmeti")
    _write_xlsx_sheet(
        ws,
        _meta_lines(user, "Potpuni izvoz - Predmeti"),
        ["Naziv", "Arhiviran"],
        [[s.name, s.is_archived] for s in Subject.objects.all()],
    )

    ws = workbook.create_sheet("Učenici")
    _write_xlsx_sheet(
        ws,
        _meta_lines(user, "Potpuni izvoz - Učenici"),
        ["Prezime", "Ime", "Razred", constants.IP_LABEL, constants.PP_LABEL, "Arhiviran"],
        [
            [s.last_name, s.first_name, s.school_class.name, s.is_ip, s.is_pp, s.is_archived]
            for s in Student.objects.select_related("school_class").order_by(
                "school_class", "last_name", "first_name"
            )
        ],
    )

    ws = workbook.create_sheet("Satovi")
    _write_xlsx_sheet(
        ws,
        _meta_lines(user, "Potpuni izvoz - Nastavni satovi"),
        ["Datum", "Razred", "Predmet", "Nastavnik", "Tema sata"],
        [
            [
                l.date.strftime("%d.%m.%Y."),
                l.school_class.name,
                l.subject.name,
                _display_name(l.teacher),
                l.topic,
            ]
            for l in Lesson.objects.select_related("school_class", "subject", "teacher").order_by(
                "date"
            )
        ],
    )

    ws = workbook.create_sheet("Prisutnost")
    _write_xlsx_sheet(
        ws,
        _meta_lines(user, "Potpuni izvoz - Prisutnost"),
        ["Datum", "Razred", "Predmet", "Učenik", "Status", "Opravdano"],
        [
            [
                a.lesson.date.strftime("%d.%m.%Y."),
                a.lesson.school_class.name,
                a.lesson.subject.name,
                a.student.full_name,
                a.status,
                "" if a.justified is None else ("da" if a.justified else "ne"),
            ]
            for a in Attendance.objects.select_related(
                "lesson__school_class", "lesson__subject", "student"
            ).order_by("lesson__date")
        ],
    )

    ws = workbook.create_sheet("Bilješke")
    header, rows = notes_data(None, None, None, None, None)
    _write_xlsx_sheet(ws, _meta_lines(user, "Potpuni izvoz - Bilješke"), header, rows)

    return _workbook_response(workbook, "potpuni_izvoz.xlsx")


# ---------------------------------------------------------------------------
# Rendering helpers
# ---------------------------------------------------------------------------


def _to_csv(meta_lines, header, rows):
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=CSV_DELIMITER)
    for line in meta_lines:
        writer.writerow([line])
    writer.writerow([])
    writer.writerow(header)
    for row in rows:
        writer.writerow(row)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _write_xlsx_sheet(ws, meta_lines, header, rows, color_status_cells=False):
    for line in meta_lines:
        ws.append([line])
    ws.append([])
    ws.append(header)
    header_row_idx = ws.max_row
    for cell in ws[header_row_idx]:
        cell.font = openpyxl.styles.Font(bold=True)
        if color_status_cells:
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="bottom")

    for row in rows:
        ws.append(row)
        if color_status_cells:
            row_idx = ws.max_row
            for col_idx in range(2, len(row) + 1):  # skip the "Učenik" column
                cell = ws.cell(row=row_idx, column=col_idx)
                fill_color, font_color = _status_fill(str(cell.value or ""))
                if fill_color:
                    cell.fill = openpyxl.styles.PatternFill(
                        start_color=fill_color, end_color=fill_color, fill_type="solid"
                    )
                    cell.font = openpyxl.styles.Font(color=font_color)

    if color_status_cells:
        ws.freeze_panes = ws.cell(row=header_row_idx + 1, column=2).coordinate
        ws.row_dimensions[header_row_idx].height = 30

    for col_idx in range(1, len(header) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = (
            22 if color_status_cells else 20
        )


def _attachment_response(content, filename, content_type):
    response = HttpResponse(content, content_type=content_type)
    ascii_fallback = filename.encode("ascii", "ignore").decode("ascii") or "izvoz"
    response["Content-Disposition"] = (
        f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quote(filename)}"
    )
    return response


def _workbook_response(workbook, filename):
    buf = io.BytesIO()
    workbook.save(buf)
    return _attachment_response(
        buf.getvalue(),
        filename,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
