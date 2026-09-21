import io

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from dnevnik import constants
from dnevnik.models import Attendance, ExportLog, Lesson, Note, SchoolClass, SchoolYear, Student, Subject


def make_user(username, **kwargs):
    user = User.objects.create_user(username=username, password="x", **kwargs)
    user.profile.must_change_password = False
    user.profile.save()
    return user


class ExportScreensTests(TestCase):
    def setUp(self):
        self.teacher1 = make_user("nastavnik1", first_name="Prvi", last_name="Nastavnik")
        self.teacher2 = make_user("nastavnik2", first_name="Drugi", last_name="Nastavnik")
        self.admin = make_user("admin1", is_staff=True)

        self.school_year = SchoolYear.objects.create(name="2025./2026.", is_active=True)
        self.school_class = SchoolClass.objects.create(name="1.a", school_year=self.school_year)
        self.subject = Subject.objects.create(name="Matematika")
        self.other_subject = Subject.objects.create(name="Hrvatski jezik")

        self.ana = Student.objects.create(
            first_name="Ana", last_name="Anić", school_class=self.school_class
        )
        self.ivan = Student.objects.create(
            first_name="Ivan", last_name="Ivić", school_class=self.school_class
        )

        self.lesson = Lesson.objects.create(
            school_class=self.school_class,
            subject=self.subject,
            teacher=self.teacher1,
            date="2026-01-10",
            topic="Uvod u razlomke",
        )
        Attendance.objects.create(
            lesson=self.lesson, student=self.ana, status=constants.ATTENDANCE_ABSENT, justified=True
        )
        Attendance.objects.create(
            lesson=self.lesson, student=self.ivan, status=constants.ATTENDANCE_PRESENT
        )

        Note.objects.create(lesson=self.lesson, student=None, author=self.teacher1, text="Opća napomena")
        Note.objects.create(lesson=self.lesson, student=self.ana, author=self.teacher1, text="Kasnila")

        self.client.login(username="nastavnik1", password="x")

    # --- Prisutnost (matrix) ---------------------------------------------

    def test_matrix_requires_razred_but_not_predmet(self):
        response = self.client.get(reverse("izvoz_prisutnost"), {"format": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/html; charset=utf-8")

    def test_matrix_xlsx_download_has_three_sheets(self):
        response = self.client.get(
            reverse("izvoz_prisutnost"),
            {"razred": self.school_class.id, "predmet": self.subject.id, "format": "xlsx"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Prisutnost", "Satovi", "Bilješke"])
        # Row 1-3 = meta lines, row 4 blank, row 5 = header, row 6+ = data.
        sheet = workbook["Prisutnost"]
        values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
        joined = " ".join(str(v) for v in values)
        self.assertIn("Anić Ana", joined)
        self.assertIn("Odsutan (opravdano)", joined)
        self.assertIn("Prisutan", joined)

    def test_matrix_bilj_sheet_contains_notes_for_same_filters(self):
        response = self.client.get(
            reverse("izvoz_prisutnost"),
            {"razred": self.school_class.id, "predmet": self.subject.id, "format": "xlsx"},
        )
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook["Bilješke"]
        values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
        joined = " ".join(str(v) for v in values)
        self.assertIn("Opća napomena", joined)
        self.assertIn("Kasnila", joined)

    def test_matrix_status_cells_are_colored(self):
        response = self.client.get(
            reverse("izvoz_prisutnost"),
            {"razred": self.school_class.id, "predmet": self.subject.id, "format": "xlsx"},
        )
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook["Prisutnost"]
        absent_cell = present_cell = None
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value == "Odsutan (opravdano)":
                    absent_cell = cell
                elif cell.value == "Prisutan":
                    present_cell = cell
        self.assertIsNotNone(absent_cell)
        self.assertIsNotNone(present_cell)
        self.assertIn("FFC7CE", str(absent_cell.fill.fgColor.rgb))
        self.assertIn("C6EFCE", str(present_cell.fill.fgColor.rgb))

    def test_matrix_all_subjects_groups_by_subject_and_date(self):
        other_lesson = Lesson.objects.create(
            school_class=self.school_class,
            subject=self.other_subject,
            teacher=self.teacher1,
            date="2026-01-11",
            topic="Padeži",
        )
        Attendance.objects.create(
            lesson=other_lesson, student=self.ana, status=constants.ATTENDANCE_LATE
        )
        response = self.client.get(
            reverse("izvoz_prisutnost"),
            {"razred": self.school_class.id, "format": "xlsx"},  # predmet left blank = svi
        )
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook["Prisutnost"]
        header_row = list(sheet.iter_rows())[4]  # 3 meta lines + 1 blank + header
        header_values = [c.value for c in header_row if c.value]
        joined = " ".join(header_values)
        self.assertIn("Matematika", joined)
        self.assertIn("Hrvatski jezik", joined)

    def test_matrix_csv_has_bom_and_semicolon_and_exporter_header(self):
        response = self.client.get(
            reverse("izvoz_prisutnost"),
            {"razred": self.school_class.id, "predmet": self.subject.id, "format": "csv"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith("﻿".encode("utf-8")))
        text = response.content.decode("utf-8-sig")
        self.assertIn(";", text)
        self.assertIn("Prvi Nastavnik", text)  # exporter name in the header

    def test_matrix_export_is_logged(self):
        self.assertEqual(ExportLog.objects.count(), 0)
        self.client.get(
            reverse("izvoz_prisutnost"),
            {"razred": self.school_class.id, "predmet": self.subject.id, "format": "xlsx"},
        )
        log = ExportLog.objects.get()
        self.assertEqual(log.report_type, "prisutnost")
        self.assertEqual(log.user, self.teacher1)

    # --- Sazetak izostanaka ------------------------------------------------

    def test_absence_summary_counts(self):
        response = self.client.get(reverse("izvoz_sazetak"), {"format": "csv"})
        text = response.content.decode("utf-8-sig")
        self.assertIn("Anić Ana;1.a;1;1;0;0", text)
        self.assertIn("Ivić Ivan;1.a;0;0;0;0", text)

    def test_absence_summary_filtered_by_teacher(self):
        response = self.client.get(
            reverse("izvoz_sazetak"), {"nastavnik": self.teacher2.id, "format": "csv"}
        )
        text = response.content.decode("utf-8-sig")
        self.assertNotIn("Anić Ana", text)

    # --- Biljeske ------------------------------------------------------

    def test_notes_export_includes_author_and_topic(self):
        response = self.client.get(reverse("izvoz_biljeske"), {"format": "csv"})
        text = response.content.decode("utf-8-sig")
        self.assertIn("Opća napomena", text)
        self.assertIn("Kasnila", text)
        self.assertIn("Prvi Nastavnik", text)
        self.assertIn("Uvod u razlomke", text)
        self.assertIn("(opća bilješka)", text)

    # --- Potpuni izvoz (admin only) ----------------------------------------

    def test_full_export_forbidden_for_regular_teacher(self):
        response = self.client.get(reverse("izvoz_potpuni"), {"preuzmi": "1"})
        self.assertEqual(response.status_code, 403)

    def test_full_export_allowed_for_admin(self):
        self.client.logout()
        self.client.login(username="admin1", password="x")
        response = self.client.get(reverse("izvoz_potpuni"), {"preuzmi": "1"})
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("Učenici", workbook.sheetnames)
        self.assertIn("Prisutnost", workbook.sheetnames)
        self.assertIn("Bilješke", workbook.sheetnames)
